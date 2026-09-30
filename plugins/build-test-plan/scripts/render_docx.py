"""Render the test-plan markdown to .docx on top of Emvigo's master template.

Usage: render_docx.py <plan.md> <plan.docx> [--spec style-spec.json] [--template master-template.docx]
                      [--assets <out>/assets] [--logo <cover logo file in assets>] [--project <name>] [--no-refresh]

The md is the source of truth. The docx starts from the master template (cover, header, footer, fonts,
heading numbering, table looks are the template's own); the template's sample body is removed and the
plan's sections are written in the same design with the template's spacing/alignment problems fixed.
Prints a JSON result with any warnings. Exit 1 on failure.
"""
import argparse
import copy
import json
import re
import subprocess
import sys
from datetime import datetime
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import qn, nsdecls
    from docx.shared import Cm, Pt, RGBColor, Twips
except ImportError:
    print(json.dumps({"ok": False, "error": "Please run `pip install python-docx`."}))
    sys.exit(1)

DEFAULT_SPEC = {
    "fonts": {"body": "Calibri", "heading": "Calibri"},
    "sizes_pt": {"body": 12, "table": 11, "table_history": 10, "h1": 14, "h2": 12, "h3": 12},
    "colors": {"placeholder": "C00000", "table_border": "000000", "history_border": "93CDDC", "history_fill": "DAEEF3"},
    "spacing_twips": {"body_after": 120, "line": 264, "h1_before": 360, "h1_after": 120, "h2_before": 240, "h2_after": 100,
                      "h3_before": 160, "h3_after": 80, "cell_before": 50, "cell_after": 50, "list_after": 60},
    "indent_twips": {"heading_text": 567, "h3_number": 567, "list_text": 567, "list_hang": 283, "list2_text": 1134},
    "tables": {"cell_margin_lr_twips": 100, "narrow_headers": ["date", "version", "rev", "revision", "s.no", "s. no", "sl no", "no", "id", "sr. no", "iteration"],
               "narrow_max_weight": 12, "min_col_weight": 8, "max_col_weight": 60, "char_twips": 105,
               "history_widths": [850, 1250, 950, 1080, 1080, 1080, 1080, 1371], "keyvalue_first_col": 3030},
    "cover": {"logo_max_width_cm": 8.0, "logo_max_height_cm": 4.0, "placeholder_text": "[Cover logo not supplied]",
              "frame_width_twips": 8741, "frame_y_twips": 12900},
    "header": {"logo_width_cm": 4.4},
    "toc": {"levels": "1-3"},
    "roman_list_headings": ["assumptions", "dependencies", "risks", "sequence & criteria for integration testing"],
}

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def merge(base, over):
    for k, v in over.items():
        if isinstance(v, dict) and isinstance(base.get(k), dict):
            merge(base[k], v)
        else:
            base[k] = v
    return base


def parse_frontmatter(text):
    fm = {}
    if text.startswith("---"):
        end = text.find("\n---", 3)
        if end != -1:
            for line in text[3:end].strip().splitlines():
                if ":" in line:
                    k, v = line.split(":", 1)
                    v = v.strip()
                    if v.startswith('"'):
                        try:
                            v = json.loads(v)
                        except ValueError:
                            pass
                    fm[k.strip()] = v
            text = text[end + 4:]
    return fm, text.lstrip("\n")


# ---------- low-level helpers ----------

def el(tag, **attrs):
    e = OxmlElement(tag if ":" in tag else "w:" + tag)
    for k, v in attrs.items():
        e.set(qn("w:" + k), str(v))
    return e


def set_run_font(run, name, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(att), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color.upper())


INLINE = re.compile(r"\*\*(.+?)\*\*|(?<!\*)\*(?!\s)(.+?)(?<!\s)\*(?!\*)|`(.+?)`")


def add_inline(paragraph, text, font, size, bold=False):
    parts = re.split(r"<br\s*/?>", text)
    for pi, part in enumerate(parts):
        if pi:
            paragraph.add_run().add_break(WD_BREAK.LINE)
        pos = 0
        for m in INLINE.finditer(part):
            if m.start() > pos:
                set_run_font(paragraph.add_run(part[pos:m.start()]), font, size, bold=bold or None)
            if m.group(1) is not None:
                set_run_font(paragraph.add_run(m.group(1)), font, size, bold=True)
            elif m.group(2) is not None:
                set_run_font(paragraph.add_run(m.group(2)), font, size, italic=True, bold=bold or None)
            else:
                set_run_font(paragraph.add_run(m.group(3)), "Consolas", size - 1)
            pos = m.end()
        if pos < len(part):
            set_run_font(paragraph.add_run(part[pos:]), font, size, bold=bold or None)


def set_ppr(p, spacing=None, ind=None, keep_next=False, page_break_before=False, num=None, jc=None):
    """Write pPr children in schema order."""
    ppr = p._p.get_or_add_pPr()
    for child in list(ppr):
        if child.tag != qn("w:pStyle"):
            ppr.remove(child)
    if keep_next:
        ppr.append(el("keepNext"))
    if page_break_before:
        ppr.append(el("pageBreakBefore"))
    if num:
        n = el("numPr")
        n.append(el("ilvl", val=num[1]))
        n.append(el("numId", val=num[0]))
        ppr.append(n)
    if spacing:
        ppr.append(el("spacing", **spacing))
    if ind:
        ppr.append(el("ind", **ind))
    if jc:
        ppr.append(el("jc", val=jc))


# ---------- template preparation ----------

def fix_styles(doc, spec):
    """Tighten the template styles: consistent spacing, keep-with-next, no hard-coded indents."""
    sp, sz = spec["spacing_twips"], spec["sizes_pt"]
    styles = doc.styles.element

    def style_all(sid):
        # the template repeats some style ids (Google Docs export); the copies must all be fixed
        return [s for s in styles.findall(qn("w:style")) if s.get(qn("w:styleId")) == sid]

    def rewrite(sid, ppr_children, size=None):
        for s in style_all(sid):
            rewrite_one(s, ppr_children, size)

    def rewrite_one(s, ppr_children, size):
        old = s.find(qn("w:pPr"))
        if old is not None:
            s.remove(old)
        ppr = el("pPr")
        for c in ppr_children:
            ppr.append(copy.deepcopy(c))
        rpr = s.find(qn("w:rPr"))
        if rpr is not None:
            s.insert(list(s).index(rpr), ppr)
        else:
            s.append(ppr)
        if size and rpr is not None:
            for t in ("w:sz", "w:szCs"):
                e = rpr.find(qn(t))
                if e is None:
                    e = OxmlElement(t)
                    rpr.append(e)
                e.set(qn("w:val"), str(int(size * 2)))

    rewrite("Normal", [el("spacing", after=sp["body_after"], line=sp["line"], lineRule="auto")])
    # Normal has no rPr in the template; size comes from docDefaults (12pt), which we keep.
    for sid, key, before, after in (("Heading1", "h1", sp["h1_before"], sp["h1_after"]),
                                    ("Heading2", "h2", sp["h2_before"], sp["h2_after"]),
                                    ("Heading3", "h3", sp["h3_before"], sp["h3_after"])):
        rewrite(sid, [el("keepNext"), el("keepLines"), el("spacing", before=before, after=after, line=sp["line"], lineRule="auto")], sz[key])

    # TOC styles: right tab at the text edge with dot leader, compact.
    tw = 8741
    for sid, left, bold in (("TOC1", 0, True), ("TOC2", 284, False), ("TOC3", 568, False)):
        for s in style_all(sid):
            old = s.find(qn("w:pPr"))
            if old is not None:
                s.remove(old)
            ppr = el("pPr")
            tabs = el("tabs")
            tabs.append(el("tab", val="left", pos=left + 567))
            tabs.append(el("tab", val="right", leader="dot", pos=tw))
            ppr.append(tabs)
            ppr.append(el("spacing", before=60, after=20))
            ppr.append(el("ind", left=left))
            rpr = s.find(qn("w:rPr"))
            if rpr is not None:
                s.remove(rpr)
            rpr = el("rPr")
            if bold:
                rpr.append(el("b"))
            rpr.append(el("sz", val=22 if bold else 21))
            rpr.append(el("szCs", val=22 if bold else 21))
            s.append(ppr)
            s.append(rpr)


def patch_numbering(doc, spec):
    """Heading numbering with aligned indents; list definitions that can be restarted per list."""
    numbering = doc.part.numbering_part.element
    ind = spec["indent_twips"]
    ht = ind["heading_text"]
    absn = None
    for a in numbering.findall(qn("w:abstractNum")):
        if a.get(qn("w:abstractNumId")) == "3":
            absn = a
    if absn is not None:
        for lvl in absn.findall(qn("w:lvl")):
            i = int(lvl.get(qn("w:ilvl")))
            if i > 2:
                continue
            ppr = lvl.find(qn("w:pPr"))
            if ppr is None:
                ppr = el("pPr")
                lvl.append(ppr)
            for t in ppr.findall(qn("w:tabs")) + ppr.findall(qn("w:ind")):
                ppr.remove(t)
            left = ht if i < 2 else ht + ind["h3_number"]
            ppr.append(el("ind", left=left, hanging=ht))
            jc = lvl.find(qn("w:lvlJc"))
            if jc is not None:
                jc.set(qn("w:val"), "left")
            suff = lvl.find(qn("w:suff"))
            if suff is not None:
                lvl.remove(suff)
    # list definitions: 91 bullet, 92 decimal, 93 lower roman
    first_num = numbering.find(qn("w:num"))

    def add_abstract(aid, fmt, text, font=None):
        a = el("abstractNum", abstractNumId=aid)
        a.append(el("multiLevelType", val="hybridMultilevel"))
        for i in range(2):
            lvl = el("lvl", ilvl=i)
            lvl.append(el("start", val=1))
            lvl.append(el("numFmt", val=fmt if i == 0 else ("bullet" if fmt == "bullet" else "lowerLetter")))
            t = text if i == 0 else ("o" if fmt == "bullet" else "%2.")
            lvl.append(el("lvlText", val=t))
            lvl.append(el("lvlJc", val="left"))
            p = el("pPr")
            left = ind["list_text"] if i == 0 else ind["list2_text"]
            p.append(el("ind", left=left, hanging=ind["list_hang"]))
            lvl.append(p)
            if fmt == "bullet":
                r = el("rPr")
                f = "Calibri" if i == 0 else "Courier New"
                r.append(el("rFonts", ascii=f, hAnsi=f, cs=f, eastAsia=f, hint="default"))
                lvl.append(r)
            a.append(lvl)
        numbering.insert(list(numbering).index(first_num), a)

    add_abstract(91, "bullet", "•")
    add_abstract(92, "decimal", "%1.")
    add_abstract(93, "lowerRoman", "%1.")


class ListNums:
    def __init__(self, doc):
        self.numbering = doc.part.numbering_part.element
        ids = [int(n.get(qn("w:numId"))) for n in self.numbering.findall(qn("w:num"))]
        self.next = max(ids + [10]) + 1

    def new(self, abstract):
        nid = self.next
        self.next += 1
        n = el("num", numId=nid)
        n.append(el("abstractNumId", val=abstract))
        o = el("lvlOverride", ilvl=0)
        o.append(el("startOverride", val=1))
        n.append(o)
        self.numbering.append(n)
        return nid


def ordinal(n):
    return "th" if 10 <= n % 100 <= 20 else {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th")


def para_text(p):
    return "".join(t.text or "" for t in p.iter(qn("w:t")))


def fill_cover(doc, fm, spec, assets_dir, warnings):
    body = doc.element.body
    paras = [p for p in body.iterchildren() if p.tag == qn("w:p")]
    cv = spec["cover"]
    project = fm.get("project", "")

    for p in paras:
        txt = para_text(p)
        if "<Insert Logo Here>" in txt:
            for r in p.findall(qn("w:r")):
                p.remove(r)
            from docx.text.paragraph import Paragraph
            para = Paragraph(p, doc._body)
            logo = fm.get("cover_logo", "N/A")
            lp = Path(assets_dir) / logo if assets_dir and logo not in ("N/A", "") else None
            if lp and lp.is_file():
                try:
                    pic = para.add_run().add_picture(str(lp), width=Cm(cv["logo_max_width_cm"]))
                    if pic.height > Cm(cv["logo_max_height_cm"]):
                        ratio = Cm(cv["logo_max_height_cm"]) / pic.height
                        pic.width, pic.height = int(pic.width * ratio), Cm(cv["logo_max_height_cm"])
                except Exception as e:
                    warnings.append(f"Cover logo could not be placed ({e}); placeholder used.")
                    lp = None
            if not (lp and lp.is_file()):
                if logo not in ("N/A", ""):
                    warnings.append(f"Cover logo file not found: {logo}")
                set_run_font(para.add_run(cv["placeholder_text"]), spec["fonts"]["heading"], 12, color=spec["colors"]["placeholder"])
        elif "<Project Name>" in txt:
            for t in p.iter(qn("w:t")):
                if t.text and "<Project Name>" in t.text:
                    t.text = t.text.replace("<Project Name>", project)
        elif "Project ID:" in txt:
            pid = fm.get("project_id", "N/A")
            if pid in ("N/A", ""):
                body.remove(p)
            else:
                for t in p.iter(qn("w:t")):
                    if t.text and "<Project ID>" in t.text:
                        t.text = t.text.replace("<Project ID>", pid)
        elif "Emvigo Technologies" in txt:
            try:
                d = datetime.fromisoformat(fm["run_at"])
            except (KeyError, ValueError):
                d = datetime.now()
            runs = p.findall(qn("w:r"))
            ts = [(r, r.find(qn("w:t"))) for r in runs if r.find(qn("w:t")) is not None]
            # runs: company (+br), day, ordinal (superscript), rest of date
            if len(ts) >= 4:
                ts[1][1].text = str(d.day)
                ts[2][1].text = ordinal(d.day)
                ts[3][1].text = d.strftime(" %B, %Y")
            else:
                warnings.append("Cover date paragraph has an unexpected shape; date left as in the template.")

    # the template's empty spacer lines keep their original height whatever Normal's spacing is
    for p in body.iterchildren():
        if p.tag == qn("w:p") and p.find(".//" + qn("w:sectPr")) is None and p.find(qn("w:pPr")) is not None:
            if not para_text(p) and not p.findall(".//" + qn("w:br")):
                ppr = p.find(qn("w:pPr"))
                for sp_ in ppr.findall(qn("w:spacing")):
                    ppr.remove(sp_)
                ppr.insert(0, el("spacing", before=0, after=0, line=240, lineRule="auto"))
    # cover -> contents: page-break-before on the contents title instead of a break paragraph that can spill
    for p in list(body.iterchildren()):
        if p.tag == qn("w:p") and para_text(p) == "Table of Contents":
            prev = p.getprevious()
            if prev is not None and prev.tag == qn("w:p") and prev.findall(".//" + qn("w:br")) and not para_text(prev):
                body.remove(prev)
            ppr = p.find(qn("w:pPr"))
            if ppr is None:
                ppr = el("pPr")
                p.insert(0, ppr)
            ppr.insert(0, el("pageBreakBefore"))
    # The company/date block is pinned to a fixed spot on page 1 (a frame), and the spacer lines that used to
    # push it down are dropped, so a wrapped project name or a tall logo can never move it onto page 2.
    kids = list(body.iterchildren())
    title_i = next(i for i, p in enumerate(kids) if para_text(p) == "for")
    comp_i = next(i for i, p in enumerate(kids) if para_text(p).startswith("Emvigo Technologies"))
    for k in kids[title_i:comp_i]:
        if k.tag == qn("w:p") and not para_text(k) and not k.findall(".//" + qn("w:br")) and not k.findall(".//" + qn("w:drawing"))                 and k.find(".//" + qn("w:sectPr")) is None:
            body.remove(k)
    comp = kids[comp_i]
    ppr = comp.find(qn("w:pPr"))
    if ppr is None:
        ppr = el("pPr")
        comp.insert(0, ppr)
    for old in ppr.findall(qn("w:framePr")):
        ppr.remove(old)
    frame = el("framePr", w=cv["frame_width_twips"], hSpace=0, wrap="around", vAnchor="page", hAnchor="margin", y=cv["frame_y_twips"])
    at = 0
    for idx, c in enumerate(ppr):
        if c.tag in (qn("w:pStyle"), qn("w:keepNext"), qn("w:keepLines"), qn("w:pageBreakBefore")):
            at = idx + 1
    ppr.insert(at, frame)


def fix_headers(doc, spec):
    """Header: smaller logo, no rule under it, no stacked empty paragraphs."""
    width_emu = int(spec.get("header", {}).get("logo_width_cm", 4.4) * 360000)
    seen = set()
    for sec in doc.sections:
        for hdr in (sec.header, sec.first_page_header):
            if hdr.is_linked_to_previous or id(hdr._element) in seen:
                continue
            seen.add(id(hdr._element))
            paras = list(hdr._element.iter(qn("w:p")))
            for p in paras:
                has_content = p.find(".//" + qn("w:drawing")) is not None or para_text(p).strip()
                if not has_content:
                    if len(paras) > 1:
                        p.getparent().remove(p)
                    continue
                ppr = p.find(qn("w:pPr"))
                if ppr is None:
                    ppr = el("pPr")
                    p.insert(0, ppr)
                for tag in ("w:pBdr", "w:spacing"):
                    for e in ppr.findall(qn(tag)):
                        ppr.remove(e)
                # schema order: pBdr, tabs, spacing, ind, jc, rPr; spacing goes before jc/rPr
                anchor = next((c for c in ppr if c.tag in (qn("w:ind"), qn("w:jc"), qn("w:rPr"))), None)
                sp_ = el("spacing", before=0, after=0)
                if anchor is not None:
                    anchor.addprevious(sp_)
                else:
                    ppr.append(sp_)
                mark = ppr.find(qn("w:rPr"))
                if mark is not None:
                    for t in ("w:sz", "w:szCs"):
                        e = mark.find(qn(t))
                        if e is not None:
                            e.set(qn("w:val"), "20")
                for r in p.findall(qn("w:r")):
                    if r.find(qn("w:drawing")) is None:
                        for tab in r.findall(qn("w:tab")):
                            r.remove(tab)  # tab stops that pushed the logo right; paragraph alignment does it now
                        continue
                    rpr = r.find(qn("w:rPr"))
                    if rpr is not None:
                        r.remove(rpr)  # the 48pt run size inflated the line height
                    for inline in r.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}inline"):
                        ext = inline.find("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent")
                        cx, cy = int(ext.get("cx")), int(ext.get("cy"))
                        ncx, ncy = width_emu, int(cy * width_emu / cx)
                        ext.set("cx", str(ncx))
                        ext.set("cy", str(ncy))
                        for a_ext in inline.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}ext"):
                            a_ext.set("cx", str(ncx))
                            a_ext.set("cy", str(ncy))
                        for k in ("distT", "distB"):
                            inline.set(k, "0")


def fix_footers(doc, fm):
    ver = fm.get("version", "N/A")
    for sec in doc.sections:
        for ftr in (sec.footer, sec.first_page_footer):
            if ftr.is_linked_to_previous:
                continue
            for p in ftr._element.iter(qn("w:p")):
                for t in p.iter(qn("w:t")):
                    s = t.text or ""
                    if s.strip() == "2":
                        t.text = ver
                    elif s.strip() == ".0":
                        t.text = ""
                    elif s.strip().startswith("Version"):
                        t.text = "Version "
                    elif s.strip() == "Page":
                        t.text = "Page "


def replace_toc(doc, spec):
    body = doc.element.body
    sdt = body.find(qn("w:sdt"))
    if sdt is None:
        return None
    content = sdt.find(qn("w:sdtContent"))
    for c in list(content):
        content.remove(c)
    p = el("p")
    ppr = el("pPr")
    ppr.append(el("spacing", before=60, after=0))
    p.append(ppr)

    def run(child):
        r = el("r")
        r.append(child)
        return r

    b = el("fldChar", fldCharType="begin")
    b.set(qn("w:dirty"), "true")
    p.append(run(b))
    it = el("instrText")
    it.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    it.text = ' TOC \\h \\z \\t "Heading 1,1,Heading 2,2,Heading 3,3" '
    p.append(run(it))
    p.append(run(el("fldChar", fldCharType="separate")))
    t = el("t")
    t.text = "Right-click here and choose Update Field to build the table of contents."
    p.append(run(t))
    p.append(run(el("fldChar", fldCharType="end")))
    content.append(p)
    return sdt


def strip_body_after_toc(doc, sdt):
    """Remove the template's sample body (after the TOC), returning the section-break paragraph."""
    body = doc.element.body
    kids = list(body.iterchildren())
    start = kids.index(sdt) + 1
    sect_p = None
    for k in kids[start:]:
        if k.tag == qn("w:sectPr"):
            continue
        if k.tag == qn("w:p") and k.find(".//" + qn("w:sectPr")) is not None:
            sect_p = k
        body.remove(k)
    if sect_p is not None:
        # plain paragraph carrying the section break (no heading style => not a TOC entry)
        ppr = sect_p.find(qn("w:pPr"))
        sp = ppr.find(qn("w:sectPr"))
        for c in list(ppr):
            ppr.remove(c)
        ppr.append(el("spacing", before=0, after=0))
        ppr.append(sp)
        for r in sect_p.findall(qn("w:r")) + sect_p.findall(qn("w:bookmarkStart")) + sect_p.findall(qn("w:bookmarkEnd")):
            sect_p.remove(r)
    return sect_p


# ---------- tables ----------

def split_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", line)]


def col_widths(rows, total, spec, size):
    t = spec["tables"]
    ncol = len(rows[0])
    header = [h.lower().strip() for h in rows[0]]
    # width needed by the longest word in each column, so nothing breaks mid-word
    need = []
    for c in range(ncol):
        words = max((len(w) for r in rows for w in re.sub(r"[*`]", "", r[c]).replace("<br>", " ").split()), default=1)
        if header[c] in t["narrow_headers"]:
            words = max(words, len(rows[0][c].strip()))  # short headers such as "S. No" stay on one line
        need.append(int(words * t["char_twips"] * size / 10) + 2 * t["cell_margin_lr_twips"] + 40)
    weights = []
    for c in range(ncol):
        longest = max(len(re.sub(r"[*`]", "", r[c])) for r in rows)
        w = max(t["min_col_weight"], min(t["max_col_weight"], longest))
        if header[c] in t["narrow_headers"]:
            w = min(w, t["narrow_max_weight"])
        weights.append(w)
    s = sum(weights)
    widths = [total * w / s for w in weights]
    # raise columns below their word minimum, taking the space from the widest ones
    for _ in range(ncol * 2):
        short = [c for c in range(ncol) if widths[c] < need[c]]
        if not short:
            break
        deficit = sum(need[c] - widths[c] for c in short)
        for c in short:
            widths[c] = need[c]
        donors = [c for c in range(ncol) if c not in short and widths[c] > need[c]]
        room = sum(widths[c] - need[c] for c in donors)
        if room <= 0:
            break
        for c in donors:
            widths[c] -= deficit * (widths[c] - need[c]) / room
    out = [int(w) for w in widths]
    out[-1] += total - sum(out)
    return out


def style_table(tbl, widths, total, border_color, border_sz, header_fill=None, first_col_bold=False, keyvalue=False):
    tblPr = tbl._tbl.tblPr
    for child in list(tblPr):
        tblPr.remove(child)
    tblPr.append(el("tblW", w=total, type="dxa"))
    tblPr.append(el("jc", val="left"))
    b = el("tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b.append(el(side, val="single", sz=border_sz, space=0, color=border_color))
    tblPr.append(b)
    tblPr.append(el("tblLayout", type="fixed"))
    grid = tbl._tbl.tblGrid
    for gc in list(grid):
        grid.remove(gc)
    for wd in widths:
        grid.append(el("gridCol", w=wd))


def build_table(doc, rows, spec, text_width, kind="grid"):
    sp, tb = spec["spacing_twips"], spec["tables"]
    ncol = max(len(r) for r in rows)
    rows = [r + [""] * (ncol - len(r)) for r in rows]
    total = text_width
    font = spec["fonts"]["body"]
    header_row = kind != "keyvalue"
    if kind == "keyvalue":
        rows = rows[1:] if rows and [c.lower() for c in rows[0][:2]] == ["field", "value"] else rows
        first = tb["keyvalue_first_col"]
        widths = [first] + [(total - first) // max(1, ncol - 1)] * (ncol - 1)
        widths[-1] += total - sum(widths)
        size = spec["sizes_pt"]["table"] + 1
    elif kind == "history":
        hw = tb["history_widths"]
        if ncol == len(hw):
            s = sum(hw)
            widths = [int(total * w / s) for w in hw]
            widths[-1] += total - sum(widths)
        else:
            widths = col_widths(rows, total, spec, spec["sizes_pt"]["table_history"])
        size = spec["sizes_pt"]["table_history"]
    else:
        size = spec["sizes_pt"]["table"]
        widths = col_widths(rows, total, spec, size)

    tbl = doc.add_table(rows=len(rows), cols=ncol)
    if kind == "history":
        style_table(tbl, widths, total, spec["colors"]["history_border"], 6)
    else:
        style_table(tbl, widths, total, spec["colors"]["table_border"], 4)

    lr = 80 if kind == "history" else tb["cell_margin_lr_twips"]
    for ri, row in enumerate(rows):
        tr = tbl.rows[ri]
        trpr = tr._tr.get_or_add_trPr()
        trpr.append(el("cantSplit"))
        if ri == 0 and header_row:
            trpr.append(el("tblHeader"))
        for ci, txt in enumerate(row):
            cell = tr.cells[ci]
            cell.width = Twips(widths[ci])
            tcpr = cell._tc.get_or_add_tcPr()
            mar = el("tcMar")
            for side, v in (("top", 0), ("left", lr), ("bottom", 0), ("right", lr)):
                mar.append(el(side, w=v, type="dxa"))
            tcpr.append(mar)
            tcpr.append(el("vAlign", val="center" if kind == "history" else "top"))
            if ri == 0 and header_row and kind == "history":
                tcpr.append(el("shd", val="clear", color="auto", fill=spec["colors"]["history_fill"]))
            p = cell.paragraphs[0]
            set_ppr(p, spacing={"before": sp["cell_before"], "after": sp["cell_after"], "line": 252, "lineRule": "auto"})
            bold = (ri == 0 and header_row) or (kind == "keyvalue" and ci == 0)
            add_inline(p, txt, font, size, bold=bold)
            if len(rows) <= 10 and ri < len(rows) - 1:
                p._p.get_or_add_pPr().insert(0, el("keepNext"))
    prev = tbl._tbl.getprevious()
    if prev is not None and prev.tag == qn("w:p") and prev.find(".//" + qn("w:numPr")) is None:
        prev_ppr = prev.get_or_add_pPr() if hasattr(prev, "get_or_add_pPr") else prev.find(qn("w:pPr"))
        if prev_ppr is not None and prev_ppr.find(qn("w:keepNext")) is None:
            prev_ppr.insert(0, el("keepNext"))
    return tbl


def spacer(doc, twips=140):
    p = doc.add_paragraph()
    set_ppr(p, spacing={"before": 0, "after": 0, "line": twips, "lineRule": "exact"})
    return p


# ---------- document ----------

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("md")
    ap.add_argument("docx")
    ap.add_argument("--spec")
    ap.add_argument("--template")
    ap.add_argument("--assets")
    ap.add_argument("--logo", default="")  # cover logo: a file name inside --assets
    ap.add_argument("--project", default="")  # overrides the project name taken from the md title
    ap.add_argument("--header-logo")  # accepted for compatibility; the template carries its own header logo
    ap.add_argument("--no-refresh", action="store_true", help="skip the Word round-trip that fills the table of contents")
    a = ap.parse_args()
    warnings = []

    spec = copy.deepcopy(DEFAULT_SPEC)
    if a.spec and Path(a.spec).is_file():
        try:
            merge(spec, json.loads(Path(a.spec).read_text(encoding="utf-8")))
        except ValueError:
            warnings.append("style-spec.json is not valid JSON; defaults used.")
    else:
        warnings.append("style-spec.json not found; defaults used.")

    tpl = Path(a.template) if a.template else Path(__file__).resolve().parent.parent / "skills" / "test-plan-template" / "assets" / "master-template.docx"
    if not tpl.is_file():
        print(json.dumps({"ok": False, "error": f"Master template not found: {tpl}"}))
        sys.exit(1)
    try:
        text = Path(a.md).read_text(encoding="utf-8")
    except OSError as e:
        print(json.dumps({"ok": False, "error": f"Cannot read markdown: {e}"}))
        sys.exit(1)
    fm, body_md = parse_frontmatter(text)
    title = fm.get("title", "")
    fm["project"] = a.project.strip() or re.sub(r"\s+Test Plan$", "", title).strip() or "Test Plan"
    fm["cover_logo"] = a.logo.strip() or fm.get("cover_logo", "N/A")
    fm["run_at"] = datetime.now().isoformat(timespec="seconds")
    fm["version"] = fm.get("version", "").strip() or "1.0"

    doc = Document(str(tpl))
    sec = doc.sections[0]
    text_width = int((sec.page_width - sec.left_margin - sec.right_margin) / 635)

    fix_styles(doc, spec)
    patch_numbering(doc, spec)
    lists = ListNums(doc)
    cp = doc.core_properties
    cp.author = ""
    cp.title = f"{fm['project']} Test Plan"
    cp.subject = fm["project"]
    cp.comments = ""
    cp.last_modified_by = ""

    fill_cover(doc, fm, spec, a.assets, warnings)
    fix_headers(doc, spec)
    fix_footers(doc, fm)
    sdt = replace_toc(doc, spec)
    if sdt is None:
        print(json.dumps({"ok": False, "error": "Master template has no table-of-contents block."}))
        sys.exit(1)
    sect_p = strip_body_after_toc(doc, sdt)

    ind, sp = spec["indent_twips"], spec["spacing_twips"]
    font, size = spec["fonts"]["body"], spec["sizes_pt"]["body"]
    hfont = spec["fonts"]["heading"]
    lines = body_md.splitlines()
    i, para = 0, []
    skipped_title = False
    counters = [0, 0, 0]
    cur_title = ""
    section_break_done = False
    first_heading = True
    body_el = doc.element.body
    final_sect = body_el.find(qn("w:sectPr"))

    def flush():
        nonlocal para
        if para:
            p = doc.add_paragraph()
            set_ppr(p)
            add_inline(p, " ".join(para), font, size)
            para = []

    def heading(n, title):
        nonlocal first_heading, section_break_done, cur_title
        lettered = re.match(r"^([A-Z])\.\s+(.*)$", title)
        numbered = re.match(r"^(\d+(?:\.\d+)*)\.?\s+(.*)$", title)
        if numbered and n == 2 and not section_break_done and sect_p is not None:
            # Introduction onwards lives in the template's second section
            body_el.insert(list(body_el).index(final_sect), sect_p)
            section_break_done = True
        level = min(n - 1, 3)
        p = doc.add_paragraph(style=f"Heading {level}")
        if lettered and level == 1:
            label = title
            set_ppr(p, keep_next=True, page_break_before=first_heading, num=(0, 0), ind={"left": 0, "firstLine": 0})
            counters[0] = counters[1] = counters[2] = 0
        else:
            label = numbered.group(2) if numbered else title
            counters[level - 1] += 1
            for k in range(level, 3):
                counters[k] = 0
            expect = ".".join(str(c) for c in counters[:level]) if level < 3 else None
            if numbered and level < 3 and numbered.group(1).rstrip(".") != expect:
                warnings.append(f"Heading number in md ({numbered.group(1)}) differs from the document numbering ({expect}): {label}")
            left = ind["heading_text"] if level < 3 else ind["heading_text"] + ind["h3_number"]
            set_ppr(p, keep_next=True, num=(3, level - 1), ind={"left": left, "hanging": ind["heading_text"]})
        first_heading = False
        cur_title = re.sub(r"^(\d+(\.\d+)*|[A-Z])\.?\s+", "", label).strip().lower()
        r = p.add_run(label)
        set_run_font(r, hfont, spec["sizes_pt"][f"h{level}"], bold=True)
        return p

    current_list = None  # (kind, numId)

    while i < len(lines):
        raw = lines[i]
        s = raw.strip()
        if not s:
            flush()
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            flush()
            current_list = None
            n, title = len(m.group(1)), m.group(2).strip()
            if n == 1 and not skipped_title:
                skipped_title = True
                i += 1
                continue
            if n == 1:
                n = 2
            heading(n, title)
            i += 1
            continue
        if s.startswith("|"):
            flush()
            current_list = None
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            rows = [split_row(l) for l in block if not re.match(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$", l)]
            if rows:
                head = cur_title
                kind = "keyvalue" if head == "document version control" else "history" if head == "document release history" else "grid"
                build_table(doc, rows, spec, text_width, kind)
                spacer(doc)
            continue
        bm = re.match(r"^(\s*)([-*+])\s+(.*)$", raw)
        nm = re.match(r"^(\s*)(\d+)[.)]\s+(.*)$", raw)
        if bm or nm:
            flush()
            mm = bm or nm
            depth = 1 if len(mm.group(1).replace("\t", "    ")) >= 2 else 0
            kind = "roman" if cur_title in spec["roman_list_headings"] else ("bullet" if bm else "decimal")
            if current_list is None or current_list[0] != kind:
                current_list = (kind, lists.new({"bullet": 91, "decimal": 92, "roman": 93}[kind]))
            p = doc.add_paragraph()
            set_ppr(p, spacing={"before": 0, "after": sp["list_after"], "line": sp["line"], "lineRule": "auto"}, num=(current_list[1], depth if kind == "bullet" else 0))
            add_inline(p, mm.group(3), font, size)
            i += 1
            continue
        if re.match(r"^-{3,}$", s):
            flush()
            i += 1
            continue
        current_list = None
        para.append(s)
        i += 1
    flush()

    if not section_break_done and sect_p is not None:
        # plan without numbered sections: the section break still has to end the first section
        body_el.insert(list(body_el).index(final_sect), sect_p)

    def save():
        try:
            doc.save(a.docx)
        except OSError as e:
            print(json.dumps({"ok": False, "error": f"Cannot write docx: {e}"}))
            sys.exit(1)

    save()
    refreshed = False
    if not a.no_refresh:
        refreshed = refresh_with_word(a.docx, warnings)
    if not refreshed:
        # Word (or the desktop app) unavailable: ask Word to build the TOC and page numbers when the file is opened
        d2 = Document(a.docx)
        st = d2.settings.element
        if st.find(qn("w:updateFields")) is None:
            st.append(el("updateFields", val="true"))
        d2.save(a.docx)
        warnings.append("Table of contents is filled when the file is opened in Word (choose Yes to update fields).")
    print(json.dumps({"ok": True, "docx": a.docx, "toc_prebuilt": refreshed, "warnings": warnings}, indent=2))


PS_REFRESH = r"""
param($path)
$ErrorActionPreference = 'Stop'
$w = New-Object -ComObject Word.Application
$w.Visible = $false
$w.DisplayAlerts = 0
try {
  $d = $w.Documents.Open($path)
  foreach ($t in $d.TablesOfContents) { $t.Update() }
  $d.Fields.Update() | Out-Null
  foreach ($t in $d.TablesOfContents) { $t.Update() }
  $d.Save()
  $d.Close($false)
} finally { $w.Quit() }
"""


def refresh_with_word(path, warnings):
    """Best effort: let Word build the table of contents and page numbers (Windows with Word installed)."""
    if sys.platform != "win32":
        return False
    try:
        script = Path(path).with_suffix(".refresh.ps1")
        script.write_text(PS_REFRESH, encoding="utf-8-sig")
        try:
            r = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(script), str(Path(path).resolve())],
                               capture_output=True, text=True, timeout=180)
        finally:
            script.unlink(missing_ok=True)
        if r.returncode != 0:
            warnings.append("Word could not pre-build the table of contents: " + (r.stderr.strip().splitlines() or ["unknown error"])[0][:160])
            return False
        return True
    except (OSError, subprocess.SubprocessError) as e:
        warnings.append(f"Word could not pre-build the table of contents: {e}")
        return False


if __name__ == "__main__":
    main()
